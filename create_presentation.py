import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    c_bg = RGBColor(23, 26, 33)        # #171a21
    c_card = RGBColor(27, 40, 56)      # #1b2838
    c_blue = RGBColor(102, 192, 244)   # #66c0f4
    c_green = RGBColor(92, 126, 16)    # #5c7e10
    c_white = RGBColor(255, 255, 255)
    c_gray = RGBColor(143, 152, 160)
    c_light_gray = RGBColor(200, 210, 220)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()
        
        # Top accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = c_blue
        bar.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, num_text):
        set_bg(slide)

        # Tag & Number
        tb_top = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_top = tb_top.text_frame.paragraphs[0]
        p_top.text = tag_text.upper()
        p_top.font.size = Pt(11)
        p_top.font.bold = True
        p_top.font.color.rgb = c_blue

        p_num = tb_top.text_frame.add_paragraph()
        p_num.text = num_text
        p_num.font.size = Pt(11)
        p_num.font.color.rgb = c_gray
        p_num.alignment = PP_ALIGN.RIGHT

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

    def add_placeholder(slide, left, top, width, height, title, sub):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(16, 18, 22)
        box.line.color.rgb = c_blue
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = "📷 " + title
        p1.alignment = PP_ALIGN.CENTER
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = c_blue

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(12)
        p2.font.color.rgb = c_gray
        return box

    # ==========================================================
    # SLIDE 1: PORTADA
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "PROYECTO INTEGRADOR — DESARROLLO FRONTEND"
    p_badge.font.size = Pt(14)
    p_badge.font.bold = True
    p_badge.font.color.rgb = c_blue
    p_badge.alignment = PP_ALIGN.CENTER

    p_title = tf.add_paragraph()
    p_title.text = "🎮 SteamLite v2.0"
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = c_white
    p_title.alignment = PP_ALIGN.CENTER

    p_sub = tf.add_paragraph()
    p_sub.text = "Plataforma SPA de Distribución Digital, Pasarela de Pago y Mini-Juegos Web"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = c_light_gray
    p_sub.alignment = PP_ALIGN.CENTER

    p_stack = tf.add_paragraph()
    p_stack.text = "\nReact 19  •  Vite 8  •  React Router v7  •  Formik + Yup  •  HTML5 Canvas 2D"
    p_stack.font.size = Pt(14)
    p_stack.font.color.rgb = c_green
    p_stack.font.bold = True
    p_stack.alignment = PP_ALIGN.CENTER

    # ==========================================================
    # SLIDE 2: DISTRIBUCIÓN POR ROLES
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Estructura del Proyecto", "Distribución por Roles y Arquitectura Modular", "02 / 07")

    roles = [
        ("👤 Integrante 1: Core Developer", 
         ["• Catálogo centralizado en games.json",
          "• Filtrado reactivo (texto, género, precio)",
          "• 3 Mini-Juegos en JS nativo y Canvas"],
         c_blue),
        ("👤 Integrante 2: UI & Vistas", 
         ["• Enrutamiento con react-router-dom v7",
          "• Componentes modulares (Navbar, Footer)",
          "• Interfaz responsiva con tema oscuro Steam"],
         c_green),
        ("👤 Integrante 3: State & Storage", 
         ["• Estado global con Context API (ShopContext)",
          "• Persistencia permanente en localStorage",
          "• Checkout con validaciones estrictas (Formik/Yup)"],
         RGBColor(255, 170, 0))
    ]

    for i, (title, points, color) in enumerate(roles):
        left = Inches(0.8 + i * 3.9)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = color
        card.line.width = Pt(2)

        ctf = card.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        for pt in points:
            p = ctf.add_paragraph()
            p.text = pt
            p.font.size = Pt(13)
            p.font.color.rgb = c_light_gray
            p.space_before = Pt(10)

    # ==========================================================
    # SLIDE 3: TIENDA Y CATÁLOGO
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Demostración Funcional", "1. Tienda y Catálogo Interactivo", "03 / 07")

    # Left text card
    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = c_card
    card.line.color.rgb = RGBColor(50, 70, 90)
    tf = card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Características del Catálogo"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue

    bullets = [
        "Carga asíncrona mediante useEffect y fetch local.",
        "Buscador en tiempo real por título y descripción.",
        "Filtro dinámico de géneros y ordenamiento por precio.",
        "Tarjetas interactivas con estados inteligentes: Comprar / En Carrito / En Biblioteca.",
        "Notificaciones flotantes (Toast) al interactuar."
    ]
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13.5)
        bp.font.color.rgb = c_light_gray
        bp.space_before = Pt(8)

    # Right Image Placeholder
    add_placeholder(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), 
                    "[Pegar Captura: Catálogo / Tienda]", 
                    "Inserta aquí la captura de pantalla de la página principal (Home.jsx)")

    # ==========================================================
    # SLIDE 4: CARRITO Y PASARELA DE PAGO
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Demostración Funcional", "2. Carrito de Compras & Checkout", "04 / 07")

    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = c_card
    card.line.color.rgb = RGBColor(50, 70, 90)
    tf = card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Flujo de Compra y Validaciones"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue

    bullets = [
        "Gestión completa: añadir items, eliminar individual y vaciar cesta.",
        "Validación estricta con Formik + Yup:",
        "  - Tarjeta: 16 dígitos numéricos exactos.",
        "  - Titular: Mínimo 3 caracteres requeridos.",
        "  - Vencimiento (MM/AA) y CVV bancario.",
        "Botón de prueba rápida '⚡ Auto-rellenar'.",
        "Modal de Confirmación estilo Steam (sustituye alert)."
    ]
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13.5)
        bp.font.color.rgb = c_light_gray
        bp.space_before = Pt(6)

    add_placeholder(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), 
                    "[Pegar Captura: Carrito y Checkout]", 
                    "Inserta aquí la captura del formulario de pago y el modal de compra exitosa")

    # ==========================================================
    # SLIDE 5: MI BIBLIOTECA Y PERSISTENCIA
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Demostración Funcional", "3. Biblioteca Personal y Persistencia", "05 / 07")

    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = c_card
    card.line.color.rgb = RGBColor(50, 70, 90)
    tf = card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Persistencia con Web Storage"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue

    bullets = [
        "Persistencia permanente con localStorage ('steamlite_library').",
        "Prevención de duplicados: se valida antes de agregar al carrito.",
        "Actualización reactiva en la tienda: cambia botón a 'En Biblioteca'.",
        "Buscador interno para filtrar colecciones grandes.",
        "Acceso directo con un clic al motor del juego ('▶ Jugar')."
    ]
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13.5)
        bp.font.color.rgb = c_light_gray
        bp.space_before = Pt(8)

    add_placeholder(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), 
                    "[Pegar Captura: Mi Biblioteca]", 
                    "Inserta aquí la captura de pantalla de la colección de juegos adquiridos")

    # ==========================================================
    # SLIDE 6: LOS 3 MINI-JUEGOS
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Motores de Juego en JavaScript", "4. Tres Mini-Juegos Integrados", "06 / 07")

    games = [
        ("🐍 Snake Retro (g1)", 
         "Motor matricial, colisiones de pared y cuerpo, récord histórico guardado en localStorage y controles WASD/flechas.",
         "[Captura Snake]"),
        ("🏓 Neon Pong (g2)", 
         "Renderizado en Canvas 2D a 60 FPS, estética synthwave, física con aceleración, IA predictiva y marcador a 5 puntos.",
         "[Captura Neon Pong]"),
        ("🚀 Space Defender (g3)", 
         "Arcade shooter con estrellas en paralaje, ráfagas láser, proyectiles alienígenas, vidas y condición de victoria.",
         "[Captura Space Defender]")
    ]

    for i, (gtitle, gdesc, gplace) in enumerate(games):
        left = Inches(0.8 + i * 3.9)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_blue
        card.line.width = Pt(1.5)

        gtf = card.text_frame
        gtf.word_wrap = True
        gp1 = gtf.paragraphs[0]
        gp1.text = gtitle
        gp1.font.size = Pt(16)
        gp1.font.bold = True
        gp1.font.color.rgb = c_white

        gp2 = gtf.add_paragraph()
        gp2.text = gdesc
        gp2.font.size = Pt(12)
        gp2.font.color.rgb = c_light_gray
        gp2.space_before = Pt(8)

        # Inner mini placeholder
        add_placeholder(s6, left + Inches(0.2), Inches(4.2), Inches(3.3), Inches(2.1), gplace, "Pega la captura del juego")

    # ==========================================================
    # SLIDE 7: CONCLUSIONES Y PREGUNTAS
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Cierre de la Exposición", "Conclusiones y Preguntas del Jurado", "07 / 07")

    card1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(3.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_card
    card1.line.color.rgb = c_green
    tf1 = card1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "✅ Objetivos Alcanzados"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = c_green

    c_points = [
        "SPA robusta con Vite y React 19.",
        "Estado global limpio sin boilerplate con Context API.",
        "Validación estricta y profesional con Formik + Yup.",
        "Persistencia confiable en localStorage.",
        "Tres motores interactivos desarrollados en JavaScript."
    ]
    for cp in c_points:
        p = tf1.add_paragraph()
        p.text = "• " + cp
        p.font.size = Pt(13)
        p.font.color.rgb = c_light_gray
        p.space_before = Pt(6)

    card2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(3.8))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_card
    card2.line.color.rgb = c_blue
    tf2 = card2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "🚀 Buenas Prácticas Aplicadas"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = c_blue

    b_points = [
        "Control de versiones Git Flow en rama feature/steamlite-v2.",
        "Arquitectura modular y componentes desacoplados.",
        "Diseño inmersivo estilo Steam y 100% responsivo.",
        "Canvas 2D optimizado con requestAnimationFrame a 60 FPS."
    ]
    for bp in b_points:
        p = tf2.add_paragraph()
        p.text = "• " + bp
        p.font.size = Pt(13)
        p.font.color.rgb = c_light_gray
        p.space_before = Pt(6)

    # Bottom bar for Q&A
    bottom_bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.8))
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = RGBColor(16, 24, 34)
    bottom_bar.line.color.rgb = c_blue
    btf = bottom_bar.text_frame
    bp = btf.paragraphs[0]
    bp.text = "🎮 ¡Muchas Gracias! — Espacio abierto para preguntas y comentarios del jurado"
    bp.font.size = Pt(16)
    bp.font.bold = True
    bp.font.color.rgb = c_white
    bp.alignment = PP_ALIGN.CENTER

    output_path = r"C:\Users\MARCOS\Desktop\steamlite\Presentacion_SteamLite.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
